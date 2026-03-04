#!/usr/bin/env python3
"""
slide_detector 결과 + (PPTX) + 전사본 → 교정 + merged JSON 생성 (1패스)

1. (PPTX 있을 때) 슬라이드 이미지 → Gemini로 PPTX 번호 식별 (캐시)
   (PPTX 없을 때) detected slide_no 직접 사용
2. 타임라인 기반 슬라이드별 세그먼트 배분
3. 슬라이드 컨텍스트로 전사 교정 (Gemini, 체크포인트 지원)
4. 교정본으로 merged JSON + transcribed.json 저장
"""

import json
import os
import re
import time
from pathlib import Path

from dotenv import load_dotenv
from google import genai
from google.genai import types

load_dotenv()

GEMINI_API_KEY = os.getenv("GOOGLE_API_KEY") or os.getenv("GEMINI_API_KEY")
GEMINI_MODEL   = "gemini-2.5-flash"
client = genai.Client(api_key=GEMINI_API_KEY)

BATCH_SIZE = 60

# ── 엔티티 후처리 필터 ──────────────────────────────────────────────
_ALPHA_DIGIT_RE    = re.compile(r'^[A-Za-z]+\d+$')
_TRAILING_LABEL_RE = re.compile(r'[\s\-][A-D]\d*$')
_TRAILING_NUM_RE   = re.compile(r'\s\d+$')
_DIAGRAM_EN_WORDS  = {"right", "left", "up", "down", "top", "bottom"}
_JOSA_SUFFIXES     = (
    "에서는", "으로는", "이라는", "에서", "에게", "에는",
    "으로", "이나", "이라", "하는",
    "을", "를", "이", "가", "은", "는", "에", "로",
    "도", "의", "와", "과", "나", "고",
)

# ASCII 고유명사를 이 카테고리로 뭉치면 안 됨
_CATEGORY_KO = {
    "레지스터", "시스템 호출", "시스템 호출 함수", "함수", "변수",
    "프로그램", "코드", "파일", "기계 명령", "인터럽트", "메모리",
    "프로세스", "스레드", "라이브러리", "하드웨어", "디스크",
}

def _is_bad_entry(k: str, v: str) -> bool:
    """정규화 맵에 포함해선 안 되는 항목 판정"""
    if k == v:
        return False
    # 1. 키 또는 값이 다이어그램 레이블 (APP2, C3 등)
    if _ALPHA_DIGIT_RE.match(k) or _ALPHA_DIGIT_RE.match(v):
        return True
    # 2. 방향 영단어 레이블
    if k.isascii() and k.replace(" ", "").isalpha() and k.lower() in _DIAGRAM_EN_WORDS:
        return True
    # 3. "단어 A/B" 레이블
    if _TRAILING_LABEL_RE.search(k):
        return True
    # 4. "단어 숫자" 레이블 (슬라이드 64, 애플리케이션 2)
    if _TRAILING_NUM_RE.search(k):
        return True
    # 5. ASCII 고유명사 → 한국어 카테고리 (RAX→레지스터, API→시스템 호출, exit→함수)
    if k.isascii() and k.strip() and v in _CATEGORY_KO:
        return True
    # 6. 조사 붙은 항목 (커널이→커널, 시스템 호출을→시스템 호출)
    for j in _JOSA_SUFFIXES:
        if k.endswith(j) and k[:-len(j)] == v:
            return True
    # 7. 과도한 일반화: 값이 키의 부분 문자열 (함수 호출→함수, C 프로그램→C, 커널 스택→스택)
    if len(v) < len(k) and v in k:
        # STT 노이즈 접두사: 값 바로 앞에 공백 없이 1~2글자 (시표준→표준)
        if k.endswith(v):
            prefix = k[: len(k) - len(v)]
            if " " not in prefix and len(prefix) <= 2:
                return False
        return True
    return False


_token_usage: dict[str, int] = {"input": 0, "output": 0, "calls": 0}


def _add_usage(response) -> None:
    u = getattr(response, "usage_metadata", None)
    if u:
        _token_usage["input"]  += getattr(u, "prompt_token_count", 0) or 0
        _token_usage["output"] += getattr(u, "candidates_token_count", 0) or 0
    _token_usage["calls"] += 1


def api_call_with_retry(func, max_retries=5, initial_wait=10):
    for attempt in range(max_retries):
        try:
            return func()
        except Exception as e:
            err = str(e)
            if any(c in err for c in ["429", "503", "500", "RESOURCE_EXHAUSTED"]) and attempt < max_retries - 1:
                wait = initial_wait * (attempt + 1)
                print(f"  재시도 ({attempt+1}): {err[:60]}, {wait}초 대기")
                time.sleep(wait)
            else:
                raise


def fmt_ts(sec: float) -> str:
    m, s = int(sec) // 60, int(sec) % 60
    return f"{m:02d}:{s:02d}"


def extract_slide_text(img_path: str) -> dict:
    """슬라이드 이미지에서 Gemini Vision으로 제목 + 본문 텍스트 추출"""
    with open(img_path, "rb") as f:
        img_bytes = f.read()

    prompt = """이 강의 슬라이드 이미지에서 텍스트를 추출하세요.

## 출력 형식 (JSON만)
{"title": "슬라이드 제목", "text": "본문 내용 전체"}

### 규칙
- 제목: 슬라이드 상단의 큰 글씨 (없으면 빈 문자열)
- 본문: 제목을 제외한 나머지 텍스트 전부 (줄바꿈은 \\n으로)
- 다이어그램/표의 텍스트도 포함
- 이미지 속 텍스트를 있는 그대로 옮길 것"""

    def call():
        return client.models.generate_content(
            model=GEMINI_MODEL,
            contents=[
                types.Part.from_bytes(data=img_bytes, mime_type="image/jpeg"),
                types.Part.from_text(text=prompt),
            ],
            config=types.GenerateContentConfig(
                temperature=0.0,
                max_output_tokens=4096,
                thinking_config=types.ThinkingConfig(thinking_budget=0),
            ),
        )

    try:
        response = api_call_with_retry(call)
        _add_usage(response)
        raw = (response.text or "").strip()
        if "```json" in raw:
            raw = raw.split("```json")[1].split("```")[0].strip()
        elif "```" in raw:
            raw = raw.split("```")[1].split("```")[0].strip()
        parsed = json.loads(raw)
        return {"title": parsed.get("title", ""), "text": parsed.get("text", "")}
    except Exception as e:
        print(f"  [텍스트 추출 오류] {e}")
        return {"title": "", "text": ""}


def load_pptx_slides(pptx_path: str) -> list[dict]:
    from pptx import Presentation
    prs = Presentation(pptx_path)
    slides = []
    for i in range(len(prs.slides)):
        slide = prs.slides[i]
        texts = [s.text.strip() for s in slide.shapes
                 if hasattr(s, "text") and s.text.strip()]
        title = texts[0] if texts else f"슬라이드 {i+1}"
        body  = "\n".join(texts[1:]) if len(texts) > 1 else ""
        slides.append({"slide_number": i + 1, "title": title, "text": body})
    return slides


def identify_pptx_slide(img_path: str, pptx_slides: list[dict]) -> int | None:
    """슬라이드 이미지를 Gemini에 보내 PPTX 슬라이드 번호 식별"""
    with open(img_path, "rb") as f:
        img_bytes = f.read()

    slides_list = "\n".join(
        f"{s['slide_number']}. {s['title']} / {s['text'][:60]}"
        for s in pptx_slides
    )

    prompt = f"""아래 슬라이드 이미지를 보고, 다음 목록에서 일치하는 슬라이드 번호를 찾아주세요.
숫자 하나만 출력하세요. 일치하는 것이 없으면 0을 출력하세요.

## 슬라이드 목록
{slides_list}

## 출력 형식
숫자만 (예: 54)
"""

    def call():
        return client.models.generate_content(
            model=GEMINI_MODEL,
            contents=[
                types.Part.from_bytes(data=img_bytes, mime_type="image/jpeg"),
                types.Part.from_text(text=prompt),
            ],
            config=types.GenerateContentConfig(
                temperature=0.0,
                max_output_tokens=1024,
                thinking_config=types.ThinkingConfig(thinking_budget=0),
            ),
        )

    response = api_call_with_retry(call)
    _add_usage(response)
    if response is None or not response.candidates:
        return None
    raw = (response.text or "").strip()
    try:
        return int(raw)
    except ValueError:
        import re
        m = re.search(r"\d+", raw)
        return int(m.group()) if m else None


def segments_in_range(segments: list[dict], start: float, end: float) -> list[dict]:
    return [s for s in segments if s["end"] > start and s["start"] < end]


def parse_batch_response(text: str) -> tuple[dict[int, str], dict[str, str]]:
    """LLM 응답에서 corrections + norm_map 동시 파싱"""
    if not text:
        return {}, {}
    raw = text.strip()
    if "```json" in raw:
        raw = raw.split("```json")[1].split("```")[0].strip()
    elif "```" in raw:
        raw = raw.split("```")[1].split("```")[0].strip()
    try:
        parsed = json.loads(raw)
    except json.JSONDecodeError:
        return {}, {}
    if not isinstance(parsed, dict):
        return {}, {}

    corrections: dict[int, str] = {}
    for c in parsed.get("corrections", []):
        if not isinstance(c, dict):
            continue
        idx = c.get("index")
        txt = c.get("text", "")
        if isinstance(idx, int) and txt:
            corrections[idx] = txt

    _STOPWORDS = {
        "시간", "개수", "주소", "전송", "데이터", "명령", "역할",
        "기능", "사용", "정의", "목적", "규칙", "방법", "결과", "내용",
    }
    norm_map: dict[str, str] = {}
    for k, v in parsed.get("norm_map", {}).items():
        if not (isinstance(k, str) and isinstance(v, str)):
            continue
        k, v = k.strip(), v.strip()
        if not k or not v:
            continue
        if len(k) <= 1:
            continue
        if k in _STOPWORDS or v in _STOPWORDS:
            continue
        if _is_bad_entry(k, v):
            continue
        norm_map[k] = v

    return corrections, norm_map


def correct_batch(
    batch: list[tuple[int, dict]],
    slide_context: str,
    existing_norm: dict[str, str],
    slide_image_path: str | None = None,
) -> tuple[dict[int, str], dict[str, str]]:
    if not batch:
        return {}, {}

    seg_text = "\n".join(
        f"[{local_i}] {seg.get('text_original', seg['text'])}"
        for local_i, (_, seg) in enumerate(batch)
    )

    has_image = slide_image_path and Path(slide_image_path).exists()
    if has_image:
        ref_block = "\n## 강의 슬라이드 이미지 (첨부됨)\n이미지에 보이는 용어, 수식, 다이어그램을 참고하여 전사를 교정하세요.\n"
        if slide_context.strip():
            ref_block += f"\n## 강의자료 텍스트 (추가 참조)\n{slide_context[:1500]}\n"
    else:
        ref_block = f"\n## 강의자료 (용어 참조)\n{slide_context[:2000]}\n" if slide_context.strip() else ""
    canonical_list = "\n".join(f'  - "{v}"' for v in sorted(set(existing_norm.values())))

    prompt = f"""강의 전사 교정 + 기술 엔티티 추출
{ref_block}
## 전사 (교정 대상)
{seg_text}

## 기존 정규형 목록 (동일 개념이면 이 중 하나로 매핑)
{canonical_list}

## 출력 (JSON만)
{{"corrections": [{{"index": 0, "text": "교정된 텍스트"}}, ...], "norm_map": {{"변형 표현": "정규형"}}}}

### 교정 규칙
- 전문용어 오타·전사 오류 수정 (강의자료의 표기를 따를 것)
- 추임새(자, 뭐, 어, 그) 제거
- 자연스러운 문장으로 다듬기
- 실제 발화 내용을 바꾸지 말 것

### 엔티티 정규화 규칙

**사용 목적**: 이 맵은 **문맥 없이** Redis 사전으로 사용된다. 어떤 텍스트에서든 키를 만나면 값으로 치환한다.
따라서 키가 **다른 의미로도 쓰일 수 있는 단어**이면 절대 포함하지 말 것.

**핵심**: 정규화란 **같은 대상의 다른 표기**를 통일하는 것이다.
- ✓ "인트랍트" → "인터럽트" ("인트랍트"는 다른 뜻이 없으므로 안전)
- ✓ "시스콜" → "시스템 호출" ("시스콜"은 다른 뜻이 없으므로 안전)
- ✗ "커널이" → "커널" (조사 붙은 형태는 정규화 대상이 아님)

**매핑 전 자문 3가지**:
1. "키와 값이 정말 같은 것의 다른 표기인가?" → NO이면 제외
2. "이 키를 다른 맥락에서 만나도 항상 이 값으로 치환해도 되는가?" → NO이면 제외
3. "이 키가 국어사전/위키사전에 독립된 뜻으로 등록된 일반 단어인가?" → YES이면 제외

**절대 금지**:
- 사전에 등록된 일반어를 기술 용어로 매핑: 그늘→커널 ✗, 큰일→커널 ✗, 리더→read ✗, 컨퍼→커널 ✗ (발음이 비슷해도 독립된 뜻이 있으면 치환 불가)
- 별개 CS 개념 통합: 트랩→인터럽트 ✗, API→시스템 호출 ✗, 버퍼→출력 버퍼 ✗
- 유형으로 묶기: RAX→레지스터 ✗, exit→함수 ✗ (고유명은 고유명 그대로)
- 하위→상위: C 프로그램→C ✗, 매개변수→변수 ✗, 커널 스택→스택 ✗
- 다이어그램 레이블, 3어절 이상 구문, 수치, 동사/형용사

**허용**: 외래어 변형, 약어/풀네임, 띄어쓰기, 한영 혼용(프린트 F→printf)
**금지 추가**: 조사 붙은 형태 (커널이, 시스템 호출을, CPU가 등 — 조사는 별도 처리, 엔티티 명사만 포함)"""

    contents = []
    if has_image:
        with open(slide_image_path, "rb") as f:
            img_bytes = f.read()
        contents.append(types.Part.from_bytes(data=img_bytes, mime_type="image/jpeg"))
    contents.append(types.Part.from_text(text=prompt))

    def call():
        return client.models.generate_content(
            model=GEMINI_MODEL,
            contents=contents,
            config=types.GenerateContentConfig(
                temperature=0.1,
                max_output_tokens=8192,
                thinking_config=types.ThinkingConfig(thinking_budget=0),
            ),
        )

    try:
        response = api_call_with_retry(call)
        _add_usage(response)
        local_corrections, norm_map = parse_batch_response(response.text or "")
    except Exception as e:
        print(f"  [배치 오류 무시] {e}")
        return {}, {}

    corrections: dict[int, str] = {}
    for local_i, txt in local_corrections.items():
        if 0 <= local_i < len(batch):
            global_i = batch[local_i][0]
            corrections[global_i] = txt
    return corrections, norm_map


def process(
    stem: str,
    pptx_path: str | None = None,
    output_dir: str = "./output",
    work_dir: str = ".",
) -> Path:
    """
    슬라이드 감지 결과 + 전사본 → 교정 + merged JSON 생성.

    Args:
        stem:       영상 파일 이름 (확장자 제외)
        pptx_path:  PPTX 파일 경로 (없으면 None)
        output_dir: slide_detector 출력 디렉토리 (output/{stem}/ 상위)
        work_dir:   transcribed.json / merged.json 위치

    Returns:
        생성된 merged JSON 경로
    """
    global _token_usage
    _token_usage = {"input": 0, "output": 0, "calls": 0}

    use_pptx  = bool(pptx_path)
    work      = Path(work_dir).resolve()
    out       = Path(output_dir).resolve()

    detector_log    = out / stem / f"{stem}_log.json"
    img_dir         = out / stem
    transcript_path = work / f"{stem}_transcribed.json"
    output_path     = work / f"{stem}_merged.json"
    checkpoint_path = work / f"{stem}_merge_checkpoint.json"

    # ── [1/4] 데이터 로드
    print("[1/4] 데이터 로드")
    with open(detector_log, encoding="utf-8") as f:
        log = json.load(f)
    with open(transcript_path, encoding="utf-8") as f:
        segments = json.load(f)

    if use_pptx:
        pptx_slides = load_pptx_slides(pptx_path)
        print(f"  PPTX 슬라이드: {len(pptx_slides)}개")
    else:
        pptx_slides = []
        print("  PPTX 없음 — detected slide_no 직접 사용")
    print(f"  전사 세그먼트: {len(segments)}개")

    saved_slides = [r for r in log["slides"] if r["status"] == "SAVED"]
    timeline     = log["timeline"]

    # ── [2/4] 슬라이드 번호 식별
    cache_path = img_dir / "slide_id_cache.json"

    if not use_pptx:
        detected_to_pptx = {row["slide_no"]: row["slide_no"] for row in saved_slides}
        print(f"\n[2/4] PPTX 없음 — slide_no 직접 매핑 ({len(detected_to_pptx)}개)")
    elif cache_path.exists():
        with open(cache_path, encoding="utf-8") as f:
            raw = json.load(f)
        detected_to_pptx = {int(k): v for k, v in raw.items()}
        print(f"\n[2/4] 슬라이드 식별 캐시 로드 ({len(detected_to_pptx)}개, Gemini 생략)")
    else:
        print(f"\n[2/4] Gemini 슬라이드 번호 식별 ({len(saved_slides)}개)")
        detected_to_pptx = {}
        for row in saved_slides:
            det_no   = row["slide_no"]
            img_path = str(img_dir / f"slide_{det_no:03d}_start.jpg")
            if not Path(img_path).exists():
                print(f"  slide_{det_no:03d}: 이미지 없음, 스킵")
                detected_to_pptx[det_no] = None
                continue
            pptx_no = identify_pptx_slide(img_path, pptx_slides)
            detected_to_pptx[det_no] = pptx_no
            title = next((s["title"] for s in pptx_slides if s["slide_number"] == pptx_no), "?")
            print(f"  slide_{det_no:03d} → PPTX {pptx_no}: {title[:40]}")
            time.sleep(0.3)
        with open(cache_path, "w", encoding="utf-8") as f:
            json.dump(detected_to_pptx, f, ensure_ascii=False)
        print(f"  캐시 저장 → {cache_path}")

    # ── [3/4] 세그먼트 배분 + 인라인 교정
    print(f"\n[3/4] 슬라이드별 세그먼트 배분 + 교정")

    slide_occurrences: dict[int, list[dict]] = {}
    slide_det_no: dict[int, int] = {}  # pptx_no → det_no 매핑 (이미지 경로용)
    for entry in timeline:
        pptx_no = detected_to_pptx.get(entry["slide_no"])
        if not pptx_no:
            continue
        slide_det_no.setdefault(pptx_no, entry["slide_no"])
        slide_occurrences.setdefault(pptx_no, []).append({
            "start_sec": entry["start_sec"],
            "end_sec":   entry["end_sec"],
            "duration":  entry["duration"],
            "is_dup":    entry["is_dup"],
        })

    # 세그먼트 → 슬라이드 매핑
    seg_slide: dict[int, int] = {}
    for slide_no, occs in slide_occurrences.items():
        for occ in occs:
            for i, seg in enumerate(segments):
                mid = (seg["start"] + seg["end"]) / 2
                if occ["start_sec"] <= mid < occ["end_sec"]:
                    seg_slide[i] = slide_no

    groups: dict[int, list[tuple[int, dict]]] = {}
    no_slide: list[tuple[int, dict]] = []
    for i, seg in enumerate(segments):
        if i in seg_slide:
            groups.setdefault(seg_slide[i], []).append((i, seg))
        else:
            no_slide.append((i, seg))

    # 기존 정규화 맵 로드
    norm_map_path = work / "normalization_map.json"
    existing_norm: dict[str, str] = {}
    if norm_map_path.exists():
        with open(norm_map_path, encoding="utf-8") as f:
            existing_norm = json.load(f)

    # 체크포인트 로드
    if checkpoint_path.exists():
        with open(checkpoint_path, encoding="utf-8") as f:
            ckpt = json.load(f)
        all_corrections: dict[int, str] = {int(k): v for k, v in ckpt["corrections"].items()}
        all_norm_entries: dict[str, str] = ckpt.get("norm_entries", {})
        done_slides: set[int] = set(ckpt["done_slides"])
        print(f"  체크포인트 재개: {len(done_slides)}개 슬라이드 완료")
    else:
        all_corrections = {}
        all_norm_entries = {}
        done_slides = set()

    def save_checkpoint():
        with open(checkpoint_path, "w", encoding="utf-8") as f:
            json.dump({
                "corrections":  all_corrections,
                "norm_entries": all_norm_entries,
                "done_slides":  list(done_slides),
            }, f, ensure_ascii=False)

    for slide_no in sorted(slide_occurrences.keys()):
        if slide_no in done_slides:
            print(f"  슬라이드 {slide_no:3d}: 스킵 (체크포인트)")
            continue
        group = groups.get(slide_no, [])
        if not group:
            done_slides.add(slide_no)
            continue
        pptx_meta = next((s for s in pptx_slides if s["slide_number"] == slide_no), {})
        context   = f"슬라이드 제목: {pptx_meta.get('title', '')}\n{pptx_meta.get('text', '')}"
        cur_norm  = {**existing_norm, **all_norm_entries}
        det_no = slide_det_no.get(slide_no)
        img_path = str(img_dir / f"slide_{det_no:03d}_end.jpg") if det_no else None
        img_flag = " +IMG" if img_path and Path(img_path).exists() else ""
        print(f"  슬라이드 {slide_no:3d} ({pptx_meta.get('title', '')[:30]:30s}): {len(group):3d}개{img_flag}", end="", flush=True)
        for b in range(0, len(group), BATCH_SIZE):
            corrections, norm_entries = correct_batch(
                group[b:b + BATCH_SIZE], context, cur_norm,
                slide_image_path=img_path,
            )
            all_corrections.update(corrections)
            all_norm_entries.update(norm_entries)
            cur_norm.update(norm_entries)
            print(".", end="", flush=True)
            time.sleep(0.3)
        done_slides.add(slide_no)
        save_checkpoint()
        print()

    if no_slide:
        print(f"  미매핑 {len(no_slide)}개 교정 중...", end="", flush=True)
        for b in range(0, len(no_slide), BATCH_SIZE):
            corrections, norm_entries = correct_batch(no_slide[b:b + BATCH_SIZE], "", {**existing_norm, **all_norm_entries})
            all_corrections.update(corrections)
            all_norm_entries.update(norm_entries)
            print(".", end="", flush=True)
            time.sleep(0.3)
        print()

    # 교정 적용
    corrected_segments = []
    for i, seg in enumerate(segments):
        s = seg.copy()
        s["text_original"] = seg.get("text_original", seg["text"])
        s["text"] = all_corrections.get(i) or s["text_original"]
        corrected_segments.append(s)

    with open(transcript_path, "w", encoding="utf-8") as f:
        json.dump(corrected_segments, f, ensure_ascii=False, indent=2)

    # 정규화 맵 저장 (기존 항목 우선)
    if all_norm_entries:
        merged_norm = {**all_norm_entries, **existing_norm}
        with open(norm_map_path, "w", encoding="utf-8") as f:
            json.dump(merged_norm, f, ensure_ascii=False, indent=2, sort_keys=True)
        added = len(merged_norm) - len(existing_norm)
        print(f"  정규화 맵: +{added}개 → {norm_map_path.name} (총 {len(merged_norm)}개)")

    checkpoint_path.unlink(missing_ok=True)

    # ── [4/4] merged JSON 구성
    print(f"\n[4/4] merged JSON 구성")
    total_duration = corrected_segments[-1]["end"] if corrected_segments else 0

    merged_slides = []
    for slide_no in sorted(slide_occurrences.keys()):
        occs      = slide_occurrences[slide_no]
        pptx_meta = next((s for s in pptx_slides if s["slide_number"] == slide_no), {})

        all_start = min(o["start_sec"] for o in occs)
        all_end   = max(o["end_sec"]   for o in occs)
        total_dur = sum(o["duration"]  for o in occs)

        all_segs = []
        for occ in occs:
            all_segs.extend(segments_in_range(corrected_segments, occ["start_sec"], occ["end_sec"]))
        seen_starts: set = set()
        unique_segs = []
        for s in sorted(all_segs, key=lambda x: x["start"]):
            if s["start"] not in seen_starts:
                seen_starts.add(s["start"])
                unique_segs.append(s)

        if use_pptx:
            title      = pptx_meta.get("title", f"슬라이드 {slide_no}")
            slide_text = pptx_meta.get("text", "")
        else:
            # PPTX 없음 → 슬라이드 이미지에서 텍스트 추출
            det_no = slide_det_no.get(slide_no, slide_no)
            img_path = str(img_dir / f"slide_{det_no:03d}_start.jpg")
            if Path(img_path).exists():
                extracted = extract_slide_text(img_path)
                title      = extracted["title"] or f"슬라이드 {slide_no}"
                slide_text = extracted["text"]
                print(f"  슬라이드 {slide_no:3d} 이미지 → 텍스트 추출: {title[:30]}")
                time.sleep(0.3)
            else:
                title      = f"슬라이드 {slide_no}"
                slide_text = ""
        time_range = f"{fmt_ts(all_start)} ~ {fmt_ts(all_end)}"
        label      = "PPTX" if use_pptx else "슬라이드"
        print(f"  {label} {slide_no}: {title[:35]:35s} {time_range}  ({len(occs)}회, {total_dur:.0f}s)")

        merged_slides.append({
            "slide_number":        slide_no,
            "title":               title,
            "time_range":          time_range,
            "time_range_seconds":  [all_start, all_end],
            "total_duration":      round(total_dur, 1),
            "occurrences":         occs,
            "slide_text":          slide_text,
            "transcript":          " ".join(s.get("text", "") for s in unique_segs),
            "transcript_segments": unique_segs,
            "segment_count":       len(unique_segs),
        })

    result = {
        "description":               "슬라이드+전사 통합 JSON (slide_detector 기반, 1패스 교정)",
        "source_slides":             pptx_path or "(없음)",
        "source_transcript":         str(transcript_path),
        "source_detector_log":       str(detector_log),
        "total_slides":              len(merged_slides),
        "total_transcript_segments": len(corrected_segments),
        "total_duration_formatted":  fmt_ts(total_duration),
        "slides": merged_slides,
    }

    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)

    print(f"\n완료! → {output_path} ({len(merged_slides)}개 슬라이드)")
    print(f"\n[토큰 사용량]")
    print(f"  API 호출 횟수 : {_token_usage['calls']}회")
    print(f"  입력 토큰     : {_token_usage['input']:,}")
    print(f"  출력 토큰     : {_token_usage['output']:,}")
    print(f"  합계          : {_token_usage['input'] + _token_usage['output']:,}")

    return Path(output_path)


def main():
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument("stem", help="영상 파일 이름 (확장자 제외, 예: 2장3절2)")
    parser.add_argument("pptx", nargs="?", default=None, help="PPTX 파일 경로 (없으면 생략)")
    args = parser.parse_args()
    process(args.stem, args.pptx)


if __name__ == "__main__":
    main()
